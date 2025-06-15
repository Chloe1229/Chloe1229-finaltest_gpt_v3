import ast
import types
from tempfile import NamedTemporaryFile
from docx import Document
from docx.shared import Pt
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from copy import deepcopy

# Utility to load selected functions from the source file without executing the
# entire Streamlit application.
def load_create_application_docx():
    path = 'step1_to_8_step8_final_.py'
    source = open(path, 'r', encoding='utf-8').read()
    tree = ast.parse(source)
    wanted = {'set_cell_font', 'clone_row', 'create_application_docx'}
    funcs = [n for n in tree.body if isinstance(n, ast.FunctionDef) and n.name in wanted]
    module = types.ModuleType('tmp')
    module.__dict__.update({'Document': Document,
                            'Pt': Pt,
                            'deepcopy': deepcopy,
                            'WD_ALIGN_VERTICAL': WD_ALIGN_VERTICAL,
                            'WD_ALIGN_PARAGRAPH': WD_ALIGN_PARAGRAPH})
    for node in funcs:
        code = compile(ast.Module([node], type_ignores=[]), path, 'exec')
        exec(code, module.__dict__)
    return module.create_application_docx


def test_create_application_docx_basic(tmp_path):
    create_docx = load_create_application_docx()

    result = {'title_text': '변경A', 'output_1_tag': 'AR'}
    requirements = {'r1': 'req1 text', 'r2': 'req2 text'}
    selections = {'test_req_r1': '충족', 'test_req_r2': '미충족'}
    output2 = ['doc1', 'doc2']

    with NamedTemporaryFile(delete=False, suffix='.docx') as tmp:
        file_path = tmp.name
    create_docx('test', result, requirements, selections, output2, file_path)

    generated = Document(file_path)
    table = generated.tables[0]

    expected_widths = [865505, 1826260, 1019810, 1152525, 1817370]
    assert [col.width for col in table.columns] == expected_widths

    assert table.cell(4, 0).text == '변경A'
    assert table.cell(4, 2).text == 'AR'
    assert table.cell(6, 0).text == 'req1 text'
    assert table.cell(6, 3).text == '○'
    assert table.cell(7, 3).text == '×'


def test_empty_page_navigation():
    step7_results = {'k': []}
    page_list = []
    for tkey, results in step7_results.items():
        if isinstance(results, dict):
            results = [results]
            step7_results[tkey] = results
        if results:
            for idx in range(len(results)):
                page_list.append((tkey, idx))
        else:
            page_list.append((tkey, None))

    session_state = {}
    if 'step8_page' not in session_state:
        session_state['step8_page'] = 0
    page = session_state['step8_page']
    total_pages = len(page_list)
    current_key, current_idx = page_list[page]

    assert total_pages == 1
    assert current_idx is None
    assert session_state['step8_page'] == 0
